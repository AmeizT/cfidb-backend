from openpyxl import Workbook
from django.http import HttpResponse
from apps.people.models import Member
from rest_framework.decorators import action
from apps.bookkeeper.models import PaymentMethod
from openpyxl.worksheet.datavalidation import DataValidation

class TitheTemplateMixin:

    @action(detail=False, methods=["get"])
    def download_tithe_template(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tithes" # type: ignore

        headers = [
            "timestamp",
            "member_name",
            "amount",
            "payment_method",
            "reference_code",
            "notes",
        ]

        ws.append(headers) # type: ignore

        # -------------------------
        # Sample Row
        # -------------------------
        ws.append([ # type: ignore
            "2026-03-01",
            "John Doe",
            100.00,
            "Bank",
            "REF123",
            "March tithe",
        ])

        # -------------------------
        # Payment Method Dropdown
        # -------------------------
        methods = [c[0] for c in PaymentMethod.choices]

        dv_methods = DataValidation(
            type="list",
            formula1=f'"{",".join(methods)}"',
            allow_blank=False
        )
        ws.add_data_validation(dv_methods) # type: ignore
        dv_methods.add("D2:D500")

        # -------------------------
        # 👥 Members Sheet (Hidden)
        # -------------------------
        members_ws = wb.create_sheet(title="Members")

        members = Member.objects.filter(assembly=request.user.church)

        members_ws.append(["ID", "Name"])

        for m in members:
            members_ws.append([m.id, m.full_name]) # type: ignore

        last_row = members_ws.max_row

        # Dropdown from Members sheet
        member_range = f"Members!$B$2:$B${last_row}"

        dv_members = DataValidation(
            type="list",
            formula1=member_range,
            allow_blank=True
        )

        ws.add_data_validation(dv_members) # type: ignore
        dv_members.add("B2:B500")

        # Hide sheet
        members_ws.sheet_state = "hidden"

        # Freeze header
        ws.freeze_panes = "A2" # type: ignore

        # -------------------------
        # Response
        # -------------------------
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=tithe_template.xlsx"

        wb.save(response)
        return response