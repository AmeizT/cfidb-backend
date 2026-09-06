from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from django.db.models import Q
from django.http import HttpResponse
from rest_framework.decorators import action

from apps.bookkeeper.models import RevenueCategory


class RevenueTemplateMixin:

    @action(detail=False, methods=["get"])
    def download_revenue_template(self, request):
        categories = list(
            RevenueCategory.objects.filter(is_active=True)
            .filter(
                Q(assembly=request.user.church, is_standard=False)
                | Q(assembly__isnull=True, is_standard=True)
            )
            .order_by("is_standard", "name")
            .values_list("name", flat=True)
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Revenue" # type: ignore

        headers = [
            "timestamp",
            "category",
            "amount",
            "notes",
        ]

        ws.append(headers) # type: ignore
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")

        # Sample
        ws.append([ # type: ignore
            "2026-03-01",
            categories[0] if categories else "",
            1500.00,
            "Sunday collection",
        ])

        categories_ws = wb.create_sheet(title="Categories")
        categories_ws.append(["Category"])
        for category in categories:
            categories_ws.append([category])

        if categories:
            category_validation = DataValidation(
                type="list",
                formula1=f"=Categories!$A$2:$A${len(categories) + 1}",
                allow_blank=False,
            )
            category_validation.error = "Choose a category from the dropdown."
            category_validation.errorTitle = "Invalid category"
            category_validation.showErrorMessage = True
            ws.add_data_validation(category_validation) # type: ignore
            category_validation.add("B2:B1000")

        categories_ws.sheet_state = "hidden"

        ws.column_dimensions["A"].width = 14 # type: ignore
        ws.column_dimensions["B"].width = 30 # type: ignore
        ws.column_dimensions["C"].width = 16 # type: ignore
        ws.column_dimensions["D"].width = 36 # type: ignore
        ws["A2"].number_format = "yyyy-mm-dd" # type: ignore
        ws["C2"].number_format = '#,##0.00' # type: ignore

        ws.freeze_panes = "A2" # type: ignore

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=revenue_template.xlsx"

        wb.save(response)
        return response
