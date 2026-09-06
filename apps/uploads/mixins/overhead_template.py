from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from django.db.models import Q
from django.http import HttpResponse
from rest_framework.decorators import action

from apps.bookkeeper.models import OverheadType


class OverheadTemplateMixin:

    @action(detail=False, methods=["get"])
    def download_overhead_template(self, request):
        overhead_types = list(
            OverheadType.objects.filter(is_active=True)
            .filter(
                Q(is_global=True, assembly__isnull=True)
                | Q(is_global=False, assembly=request.user.church)
            )
            .order_by("is_global", "name")
            .values_list("name", flat=True)
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Overheads" # type: ignore

        headers = [
            "timestamp",
            "overhead_type",
            "amount",
            "notes",
        ]

        ws.append(headers) # type: ignore
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")

        ws.append([ # type: ignore
            "2026-03-01",
            overhead_types[0] if overhead_types else "",
            800.00,
            "Monthly rent",
        ])

        types_ws = wb.create_sheet(title="Overhead Types")
        types_ws.append(["Overhead type"])
        for overhead_type in overhead_types:
            types_ws.append([overhead_type])

        if overhead_types:
            type_validation = DataValidation(
                type="list",
                formula1=f"='Overhead Types'!$A$2:$A${len(overhead_types) + 1}",
                allow_blank=False,
            )
            type_validation.error = "Choose an overhead type from the dropdown."
            type_validation.errorTitle = "Invalid overhead type"
            type_validation.showErrorMessage = True
            ws.add_data_validation(type_validation) # type: ignore
            type_validation.add("B2:B1000")

        types_ws.sheet_state = "hidden"

        ws.column_dimensions["A"].width = 14 # type: ignore
        ws.column_dimensions["B"].width = 30 # type: ignore
        ws.column_dimensions["C"].width = 16 # type: ignore
        ws.column_dimensions["D"].width = 36 # type: ignore
        ws["A2"].number_format = "yyyy-mm-dd" # type: ignore
        ws["C2"].number_format = '#,##0.00' # type: ignore

        ws.freeze_panes = "A2" # type: ignore

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=overhead_template.xlsx"

        wb.save(response)
        return response
