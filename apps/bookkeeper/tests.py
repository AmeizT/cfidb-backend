from datetime import date
from decimal import Decimal
from io import BytesIO
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.bookkeeper.models import Expenditure, Overhead, OverheadType, Revenue, RevenueCategory, Tithe
from apps.churches.models import Church
from apps.people.models import Attendance, Member
from apps.reports.models import AssemblyReport
from apps.reports.models import AuditLog
from apps.reports.models import ReportSectionStatus
from apps.reports.services.lifecycle import get_section_source
from apps.users.models import User
from apps.bookkeeper.category_matching import normalize_financial_category_name


class ManualEntryBatchApiTests(TestCase):
    def setUp(self):
        self.assembly = Church.objects.create(name="Batch Assembly")
        self.other_assembly = Church.objects.create(name="Other Assembly")
        self.user = User.objects.create_user(
            first_name="Batch", last_name="User",
            username="batch-user", email="batch@example.com", password="password",
            church=self.assembly,
        )
        self.user.is_admin = True
        self.user.save(update_fields=["is_admin"])
        self.member = Member.objects.create(
            assembly=self.assembly, first_name="Batch", last_name="Member",
            date_of_birth=date(1990, 1, 1), gender="Male", country="Botswana",
        )
        self.other_member = Member.objects.create(
            assembly=self.other_assembly, first_name="Other", last_name="Member",
            date_of_birth=date(1990, 1, 1), gender="Male", country="Botswana",
        )
        self.report = AssemblyReport.objects.create(
            assembly=self.assembly, period_start=date(2026, 7, 1), period_end=date(2026, 7, 31),
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def envelope(self, entries):
        return {"period": "2026-07", "report": self.report.pk, "entries": entries}

    def test_tithes_batch_creates_multiple_records_and_updates_report(self):
        response = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": self.member.pk, "amount": "100.00", "payment_method": "Cash", "timestamp": "2026-07-05"},
            {"member": None, "amount": "50.00", "payment_method": "Bank", "timestamp": "2026-07-06"},
            {"member": None, "amount": "25.00", "payment_method": "Cash", "timestamp": "2026-07-07"},
        ]), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Tithe.objects.filter(report=self.report).count(), 3)
        self.assertTrue(AuditLog.objects.filter(object_id__in=Tithe.objects.filter(report=self.report).values("id"), user=self.user).exists())
        self.report.refresh_from_db()
        self.assertEqual(self.report.tithe_total, Decimal("175.00"))
        duplicate = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": self.member.pk, "amount": "10.00", "payment_method": "Cash", "timestamp": "2026-07-08"},
        ]), format="json")
        self.assertEqual(duplicate.status_code, 400)
        self.assertIn("already been recorded", str(duplicate.data))

    def test_tithes_batch_is_atomic_and_rejects_duplicates_and_cross_assembly_member(self):
        response = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": self.member.pk, "amount": "100.00", "payment_method": "Cash", "timestamp": "2026-07-05"},
            {"member": self.member.pk, "amount": "20.00", "payment_method": "Bank", "timestamp": "2026-07-06"},
        ]), format="json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Tithe.objects.filter(report=self.report).count(), 0)
        cross_scope = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": self.other_member.pk, "amount": "10.00", "payment_method": "Cash", "timestamp": "2026-07-05"},
        ]), format="json")
        self.assertEqual(cross_scope.status_code, 400)
        self.assertEqual(Tithe.objects.filter(report=self.report).count(), 0)

    def test_revenue_batch_validates_category_scope_and_uniqueness(self):
        category = RevenueCategory.objects.create(assembly=self.assembly, name="Offering")
        other = RevenueCategory.objects.create(assembly=self.other_assembly, name="Other")
        valid = self.client.post("/api/v1/bookkeeper/revenue/batch/", self.envelope([
            {"category": category.pk, "amount": "200.00", "timestamp": "2026-07-05"},
        ]), format="json")
        self.assertEqual(valid.status_code, 201, valid.data)
        duplicate = self.client.post("/api/v1/bookkeeper/revenue/batch/", self.envelope([
            {"category": category.pk, "amount": "20.00", "timestamp": "2026-07-06"},
        ]), format="json")
        self.assertEqual(duplicate.status_code, 400)
        wrong_scope = self.client.post("/api/v1/bookkeeper/revenue/batch/", self.envelope([
            {"category": other.pk, "amount": "20.00", "timestamp": "2026-07-06"},
        ]), format="json")
        self.assertEqual(wrong_scope.status_code, 400)
        self.assertEqual(Revenue.objects.filter(report=self.report).count(), 1)

    def test_overhead_batch_accepts_global_and_assembly_types(self):
        global_type = OverheadType.objects.create(name="Global Rent", is_global=True)
        local_type = OverheadType.objects.create(name="Local Power", assembly=self.assembly)
        response = self.client.post("/api/v1/bookkeeper/overhead/batch/", self.envelope([
            {"overhead_type": global_type.pk, "amount": "300.00", "timestamp": "2026-07-05"},
            {"overhead_type": local_type.pk, "amount": "80.00", "timestamp": "2026-07-05"},
        ]), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Overhead.objects.filter(report=self.report).count(), 2)

    def test_revenue_template_uses_current_scoped_categories_for_dropdown(self):
        RevenueCategory.objects.create(
            assembly=self.assembly, name="Youth Fund", needs_review=True,
        )
        RevenueCategory.objects.create(
            assembly=self.other_assembly, name="Other Assembly Income", needs_review=True,
        )
        RevenueCategory.objects.create(
            assembly=self.assembly, name="Inactive Income", needs_review=True,
            is_active=False,
        )

        response = self.client.get(
            "/api/v1/bookkeeper/revenue/download_revenue_template/"
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook["Revenue"]
        options_sheet = workbook["Categories"]
        options = [cell.value for cell in options_sheet["A"][1:]]
        validation = list(worksheet.data_validations.dataValidation)[0]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["timestamp", "category", "amount", "notes"],
        )
        self.assertEqual(options_sheet.sheet_state, "hidden")
        self.assertIn("Youth Fund", options)
        self.assertNotIn("Other Assembly Income", options)
        self.assertNotIn("Inactive Income", options)
        self.assertEqual(validation.type, "list")
        self.assertEqual(str(validation.sqref), "B2:B1000")
        self.assertEqual(
            validation.formula1,
            f"=Categories!$A$2:$A${len(options) + 1}",
        )

    def test_overhead_template_uses_current_scoped_types_for_dropdown(self):
        OverheadType.objects.create(
            assembly=self.assembly, name="Local Security", needs_review=True,
        )
        OverheadType.objects.create(
            assembly=self.other_assembly, name="Other Assembly Cost", needs_review=True,
        )
        OverheadType.objects.create(
            assembly=self.assembly, name="Inactive Cost", needs_review=True,
            is_active=False,
        )

        response = self.client.get(
            "/api/v1/bookkeeper/overhead/download_overhead_template/"
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook["Overheads"]
        options_sheet = workbook["Overhead Types"]
        options = [cell.value for cell in options_sheet["A"][1:]]
        validation = list(worksheet.data_validations.dataValidation)[0]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["timestamp", "overhead_type", "amount", "notes"],
        )
        self.assertEqual(options_sheet.sheet_state, "hidden")
        self.assertIn("Local Security", options)
        self.assertNotIn("Other Assembly Cost", options)
        self.assertNotIn("Inactive Cost", options)
        self.assertEqual(validation.type, "list")
        self.assertEqual(str(validation.sqref), "B2:B1000")
        self.assertEqual(
            validation.formula1,
            f"='Overhead Types'!$A$2:$A${len(options) + 1}",
        )

    def test_expense_batch_preserves_file_and_rolls_back_invalid_row(self):
        receipt = SimpleUploadedFile("receipt.txt", b"receipt", content_type="text/plain")
        response = self.client.post("/api/v1/bookkeeper/expenditure/batch/", {
            "period": "2026-07",
            "report": str(self.report.pk),
            "entries": json.dumps([{"name": "Fuel", "category": "travel", "quantity": 2, "price": "10.00", "invoice_date": "2026-07-05"}]),
            "entries.0.receipt": receipt,
        }, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        expense = Expenditure.objects.get(report=self.report)
        self.assertTrue(expense.receipt.name)
        self.assertTrue(expense.receipt.storage.exists(expense.receipt.name))
        payload = self.envelope([
            {"name": "Valid", "category": "travel", "quantity": 2, "price": "10.00", "invoice_date": "2026-07-05"},
            {"name": "Invalid", "category": "travel", "quantity": 0, "price": "10.00", "invoice_date": "2026-07-05"},
        ])
        response = self.client.post("/api/v1/bookkeeper/expenditure/batch/", payload, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Expenditure.objects.filter(report=self.report).count(), 1)

    def test_attendance_batch_creates_then_updates_week(self):
        created = self.client.post("/api/v1/people/attendance/batch/", self.envelope([
            {"timestamp": "2026-07-05", "service_type": "sunday", "men": 10, "women": 12},
            {"timestamp": "2026-07-12", "service_type": "sunday", "men": 8, "women": 9},
        ]), format="json")
        self.assertEqual(created.status_code, 201, created.data)
        record = Attendance.objects.get(report=self.report, timestamp=date(2026, 7, 5))
        updated = self.client.post("/api/v1/people/attendance/batch/", self.envelope([
            {"id": record.pk, "timestamp": "2026-07-05", "service_type": "sunday", "men": 15, "women": 12},
        ]), format="json")
        self.assertEqual(updated.status_code, 201, updated.data)
        record.refresh_from_db()
        self.assertEqual(record.total_adults, 27)

    def test_attendance_batch_persists_details_without_losing_metrics(self):
        created = self.client.post("/api/v1/people/attendance/batch/", self.envelope([
            {
                "timestamp": "2026-07-05", "service_type": "sunday",
                "men": 10, "women": 12, "visitor_men": 3,
            },
        ]), format="json")
        self.assertEqual(created.status_code, 201, created.data)
        record = Attendance.objects.get(report=self.report, timestamp=date(2026, 7, 5))

        updated = self.client.post("/api/v1/people/attendance/batch/", self.envelope([
            {
                "id": record.pk, "timestamp": "2026-07-05",
                "service_type": "sunday", "men": 10, "women": 12,
                "visitor_men": 3, "is_special_event": True,
                "special_event_name": "Family Sunday", "preacher": "E. Zhuwao",
                "sermon": "Faith in action", "scriptures": "James 2:14-26",
                "weather": "sunny", "notes": "Full morning service",
            },
        ]), format="json")
        self.assertEqual(updated.status_code, 201, updated.data)

        record.refresh_from_db()
        self.assertEqual(record.men, 10)
        self.assertEqual(record.women, 12)
        self.assertEqual(record.visitor_men, 3)
        self.assertEqual(record.preacher, "E. Zhuwao")
        self.assertEqual(record.sermon, "Faith in action")
        self.assertEqual(record.scriptures, "James 2:14-26")
        self.assertEqual(record.weather, "sunny")
        self.assertEqual(record.special_event_name, "Family Sunday")
        self.assertEqual(record.notes, "Full morning service")

    def test_batch_requires_an_active_assembly(self):
        user = User.objects.create_user(
            first_name="No", last_name="Assembly", username="no-assembly",
            email="no-assembly@example.com", password="password", church=None,
        )
        self.client.force_authenticate(user)
        response = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": None, "amount": "10.00", "payment_method": "Cash", "timestamp": "2026-07-05"},
        ]), format="json")
        self.assertEqual(response.status_code, 403)

    def test_invalid_batch_does_not_leave_a_new_report_behind(self):
        response = self.client.post("/api/v1/bookkeeper/tithes/batch/", {
            "period": "2026-08",
            "entries": [
                {"member": self.member.pk, "amount": "10.00", "payment_method": "Cash", "timestamp": "2026-08-02"},
                {"member": self.member.pk, "amount": "20.00", "payment_method": "Bank", "timestamp": "2026-08-09"},
            ],
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(AssemblyReport.objects.filter(
            assembly=self.assembly,
            period_start=date(2026, 8, 1),
        ).exists())

    def test_batch_requires_authentication(self):
        self.client.force_authenticate(user=None)
        response = self.client.post("/api/v1/bookkeeper/tithes/batch/", self.envelope([
            {"member": None, "amount": "10.00", "payment_method": "Cash", "timestamp": "2026-07-05"},
        ]), format="json")
        self.assertIn(response.status_code, (401, 403))
        self.assertEqual(Tithe.objects.filter(report=self.report).count(), 0)


class FinancialCategorySuggestionApiTests(TestCase):
    def setUp(self):
        self.assembly = Church.objects.create(name="Category Assembly")
        self.other_assembly = Church.objects.create(name="Other Category Assembly")
        self.user = User.objects.create_user(
            first_name="Category", last_name="Manager",
            username="category-manager", email="categories@example.com",
            password="password", church=self.assembly,
        )
        self.user.is_admin = True
        self.user.save(update_fields=["is_admin"])
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.revenue_standard = RevenueCategory.objects.get(
            normalized_name="donation gift", is_standard=True,
        )
        self.overhead_standard = OverheadType.objects.get(
            normalized_name="facility utility", is_global=True,
        )

    def test_normalization_handles_spacing_punctuation_and_plural_forms(self):
        values = [
            "Building repair", "Building Repairs", "building-repairs",
            " Building   Repairs ",
        ]
        self.assertEqual(
            {normalize_financial_category_name(value) for value in values},
            {"building repair"},
        )

    def test_normalized_duplicate_is_prevented(self):
        RevenueCategory.objects.create(
            assembly=self.assembly, name="Building Repairs",
            standard_category=self.revenue_standard,
        )
        response = self.client.post(
            "/api/v1/bookkeeper/revenue/categories/",
            {
                "name": "building-repairs",
                "standard_category_id": self.revenue_standard.pk,
                "needs_review": False,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(RevenueCategory.objects.filter(assembly=self.assembly).count(), 1)

    def test_suggestions_are_grouped_and_isolated_by_assembly(self):
        local = OverheadType.objects.create(
            assembly=self.assembly, name="Building Repairs",
            standard_category=self.overhead_standard,
        )
        OverheadType.objects.create(
            assembly=self.other_assembly, name="Building Repair Team",
            standard_category=self.overhead_standard,
        )
        response = self.client.get(
            "/api/v1/bookkeeper/overhead/types/suggestions/",
            {"q": "building-repairs"},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["assembly_name"], self.assembly.name)
        self.assertEqual(response.data["assembly_matches"][0]["id"], local.pk)
        self.assertNotIn(
            "Building Repair Team",
            [item["name"] for item in response.data["assembly_matches"]],
        )
        self.assertIn(
            "Maintenance & Repairs",
            [item["name"] for item in response.data["standard_matches"]],
        )

    def test_custom_category_maps_to_standard_category(self):
        response = self.client.post(
            "/api/v1/bookkeeper/overhead/types/",
            {
                "name": "Building Repairs",
                "standard_category_id": self.overhead_standard.pk,
                "needs_review": False,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        category = OverheadType.objects.get(pk=response.data["id"])
        self.assertEqual(category.assembly, self.assembly)
        self.assertEqual(category.standard_category, self.overhead_standard)
        self.assertEqual(category.created_by, self.user)
        self.assertFalse(category.needs_review)
        self.assertEqual(response.data["standard_category"]["name"], "Facilities & Utilities")

        report = AssemblyReport.objects.create(
            assembly=self.assembly,
            period_start=date(2026, 7, 1),
            period_end=date(2026, 7, 31),
        )
        Overhead.objects.create(
            assembly=self.assembly,
            report=report,
            overhead_type=category,
            amount=Decimal("80.00"),
            timestamp=date(2026, 7, 8),
        )
        source = get_section_source(
            report, ReportSectionStatus.Section.OPERATING_EXPENSES,
        )
        row = source["breakdown"][0]
        self.assertEqual(row["overhead_type__name"], "Building Repairs")
        self.assertEqual(row["reporting_category"], "Facilities & Utilities")
        self.assertEqual(row["reporting_category_id"], self.overhead_standard.pk)

    def test_unmapped_category_is_marked_for_review(self):
        response = self.client.post(
            "/api/v1/bookkeeper/revenue/categories/",
            {
                "name": "Special Community Proceeds",
                "standard_category_id": None,
                "needs_review": True,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        category = RevenueCategory.objects.get(pk=response.data["id"])
        self.assertTrue(category.needs_review)
        self.assertIsNone(category.standard_category)

    def test_standard_category_can_be_selected_directly_in_entry_batch(self):
        report = AssemblyReport.objects.create(
            assembly=self.assembly,
            period_start=date(2026, 7, 1), period_end=date(2026, 7, 31),
        )
        response = self.client.post(
            "/api/v1/bookkeeper/revenue/batch/",
            {
                "period": "2026-07", "report": report.pk,
                "entries": [{
                    "category": self.revenue_standard.pk,
                    "amount": "25.00", "timestamp": "2026-07-05",
                }],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Revenue.objects.get(report=report).category, self.revenue_standard)
