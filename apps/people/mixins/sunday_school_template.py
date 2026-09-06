from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.decorators import action

from apps.people.choices.services import SundaySchoolClassChoices
from apps.people.models import Member


class SundaySchoolTemplateMixin:
    @action(detail=False, methods=["get"])
    def download_sunday_school_template(self, request):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sunday School"

        teachers = list(
            Member.objects.filter(
                assembly=request.user.church,
            ).order_by("first_name", "last_name")
        )

        headers = [
            "service_date",
            "class_name",
            "teacher_name",
            "boys",
            "girls",
            "male_visitors",
            "female_visitors",
            "male_first_timers",
            "female_first_timers",
            "lesson_title",
            "scripture_reference",
            "offering",
            "remarks",
        ]
        worksheet.append(headers)
        worksheet.append([
            "2026-09-06",
            "beginners",
            teachers[0].full_name if teachers else "",
            12,
            14,
            2,
            1,
            1,
            0,
            "God created the world",
            "Genesis 1:1",
            120.00,
            "",
        ])

        header_fill = PatternFill("solid", fgColor="E8F1FF")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:M{worksheet.max_row}"
        worksheet.column_dimensions["A"].width = 14
        worksheet.column_dimensions["B"].width = 18
        worksheet.column_dimensions["C"].width = 28
        for column in ["D", "E", "F", "G", "H", "I"]:
            worksheet.column_dimensions[column].width = 20
        worksheet.column_dimensions["J"].width = 30
        worksheet.column_dimensions["K"].width = 22
        worksheet.column_dimensions["L"].width = 14
        worksheet.column_dimensions["M"].width = 34
        for row in range(2, 501):
            worksheet.cell(row=row, column=1).number_format = "yyyy-mm-dd"
            worksheet.cell(row=row, column=12).number_format = "0.00"

        class_values = ",".join(value for value, _ in SundaySchoolClassChoices.choices)
        class_validation = DataValidation(
            type="list",
            formula1=f'"{class_values}"',
            allow_blank=False,
        )
        worksheet.add_data_validation(class_validation)
        class_validation.add("B2:B500")

        teachers_sheet = workbook.create_sheet(title="Teachers")
        teachers_sheet.append(["ID", "Name"])
        for teacher in teachers:
            teachers_sheet.append([teacher.id, teacher.full_name])

        if teachers_sheet.max_row > 1:
            teacher_validation = DataValidation(
                type="list",
                formula1=f"=Teachers!$B$2:$B${teachers_sheet.max_row}",
                allow_blank=False,
            )
            worksheet.add_data_validation(teacher_validation)
            teacher_validation.add("C2:C500")
        teachers_sheet.sheet_state = "hidden"

        whole_number_validation = DataValidation(
            type="whole",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=False,
        )
        worksheet.add_data_validation(whole_number_validation)
        whole_number_validation.add("D2:I500")

        offering_validation = DataValidation(
            type="decimal",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=True,
        )
        worksheet.add_data_validation(offering_validation)
        offering_validation.add("L2:L500")

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            'attachment; filename="sunday_school_attendance_template.xlsx"'
        )
        workbook.save(response)
        return response
