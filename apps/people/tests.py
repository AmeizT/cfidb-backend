from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.apps import apps as django_apps
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIClient
from django.urls import reverse
from openpyxl import load_workbook

from apps.bookkeeper.models import Tithe
from apps.churches.models import Church, Region, RegionLeadership, Zone
from apps.people.models import (
    AssemblyMembership,
    AssemblyMembershipStatus,
    Member,
    MemberTransferRequest,
    MembershipEndReason,
    FormerMember,
    Household,
    HouseholdMember,
    HouseholdRole,
    SundaySchoolAttendance,
)
from apps.people.services.member_transfer_service import (
    accept_transfer_request,
    cancel_transfer_request,
    create_transfer_request,
    reject_transfer_request,
    complete_transfer_request,
)
from apps.users.models import User


class SundaySchoolTemplateUploadTests(TestCase):
    template_url = (
        "/api/v1/people/sunday-school-attendance/"
        "download_sunday_school_template/"
    )
    upload_url = "/api/v1/people/sunday-school-attendance/upload_excel/"
    expected_headers = [
        "service_date",
        "class_name",
        "teacher_name",
        "boys",
        "girls",
        "male_visitors",
        "female_visitors",
        "male_first_timers",
        "female_first_timers",
        "lesson_title",
        "scripture_reference",
        "offering",
        "remarks",
    ]

    def setUp(self):
        self.assembly = Church.objects.create(name="Sunday School Upload Assembly")
        self.user = User.objects.create_user(
            first_name="Upload",
            last_name="Admin",
            username="sunday-school-upload-admin",
            email="sunday-school-upload@example.com",
            password="password",
            church=self.assembly,
        )
        self.teacher = Member.objects.create(
            assembly=self.assembly,
            first_name="Sunday",
            last_name="Teacher",
            date_of_birth=date(1990, 1, 1),
            gender="Female",
            country="Botswana",
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def download_workbook(self):
        response = self.client.get(self.template_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            "sunday_school_attendance_template.xlsx",
            response["Content-Disposition"],
        )
        return load_workbook(BytesIO(response.content))

    def test_template_headers_match_upload_schema(self):
        workbook = self.download_workbook()
        worksheet = workbook["Sunday School"]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            self.expected_headers,
        )
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(workbook["Teachers"].sheet_state, "hidden")

    def test_downloaded_template_can_be_filled_and_uploaded(self):
        workbook = self.download_workbook()
        worksheet = workbook["Sunday School"]
        worksheet["C2"] = self.teacher.full_name

        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "sunday_school_attendance.xlsx",
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        response = self.client.post(
            self.upload_url,
            {"file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["created"], 1)
        self.assertEqual(response.data["updated"], 0)
        self.assertEqual(response.data["errors"], [])
        record = SundaySchoolAttendance.objects.get(
            assembly=self.assembly,
            service_date=date(2026, 9, 6),
            class_name="beginners",
        )
        self.assertEqual(record.teacher, self.teacher)
        self.assertEqual(record.boys, 12)
        self.assertEqual(record.offering, Decimal("120.00"))


class MemberResponseContractTests(TestCase):
    def setUp(self):
        self.assembly = Church.objects.create(name="Member Contract Assembly")
        self.user = User.objects.create_user(
            first_name="Member",
            last_name="Contract Admin",
            username="member-contract-admin",
            email="member-contract@example.com",
            password="password",
            church=self.assembly,
        )
        self.member = Member.objects.create(
            assembly=self.assembly,
            first_name="Contract",
            last_name="Member",
            date_of_birth=date(1990, 1, 1),
            gender="Male",
            country="Botswana",
            membership_stage="associate",
            date_of_death=None,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_member_list_exposes_membership_stage_and_nullable_date_of_death(self):
        response = self.client.get(
            reverse("people:members-list"),
            {"page_size": 100},
        )

        self.assertEqual(response.status_code, 200)

        record = next(
            item
            for item in response.data["results"]
            if item["id"] == self.member.id
        )

        self.assertEqual(record["membership_stage"], "associate")
        self.assertIsNone(record["date_of_death"])

    def test_member_detail_serializes_date_of_death_as_iso_date(self):
        self.member.date_of_death = date(2025, 4, 3)
        self.member.save(update_fields=["date_of_death"])

        response = self.client.get(
            reverse(
                "people:members-detail",
                kwargs={"member_key": self.member.member_key},
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["membership_stage"], "associate")
        self.assertEqual(response.data["date_of_death"], "2025-04-03")


class MemberTransferWorkflowTests(TestCase):
    def setUp(self):
        self.region = Region.objects.create(name="Transfer Region", code="TR")
        self.other_region = Region.objects.create(name="Other Transfer Region", code="OTR")
        self.zone = Zone.objects.create(region=self.region, name="Transfer Zone")
        self.other_zone = Zone.objects.create(region=self.other_region, name="Other Transfer Zone")

        self.from_assembly = Church.objects.create(name="Assembly A", zone=self.zone)
        self.to_assembly = Church.objects.create(name="Assembly B", zone=self.zone)
        self.other_assembly = Church.objects.create(name="Assembly C", zone=self.other_zone)
        self.external_assembly = Church.objects.create(name="Assembly D", zone=self.other_zone)

        self.source_user = User.objects.create_user(
            first_name="Source",
            last_name="Admin",
            email="source@example.com",
            username="source-admin",
            password="password",
            church=self.from_assembly,
        )
        self.receiving_user = User.objects.create_user(
            first_name="Receiving",
            last_name="Admin",
            email="receiving@example.com",
            username="receiving-admin",
            password="password",
            church=self.to_assembly,
        )
        self.other_user = User.objects.create_user(
            first_name="Other",
            last_name="Admin",
            email="other@example.com",
            username="other-admin",
            password="password",
            church=self.other_assembly,
        )
        self.regional_user = User.objects.create_user(
            first_name="Regional",
            last_name="Staff",
            email="regional@example.com",
            username="regional-staff",
            password="password",
        )
        RegionLeadership.objects.create(
            region=self.region,
            user=self.regional_user,
            role=RegionLeadership.Role.OVERSEER,
            is_active=True,
        )
        self.member = Member.objects.create(
            assembly=self.from_assembly,
            first_name="Transfer",
            last_name="Member",
            date_of_birth=date(1990, 1, 1),
            gender="Male",
            country="Zimbabwe",
            membersince=date(2020, 1, 1),
        )

    def create_member(self, assembly, first_name, last_name="Member"):
        return Member.objects.create(
            assembly=assembly,
            first_name=first_name,
            last_name=last_name,
            date_of_birth=date(1990, 1, 1),
            gender="Male",
            country="Zimbabwe",
            membersince=date(2020, 1, 1),
        )

    def create_pending_transfer(self):
        return create_transfer_request(
            member=self.member,
            to_assembly=self.to_assembly,
            effective_date=date(2026, 7, 10),
            reason="Member relocated",
            notes="Initial request",
            requested_by=self.source_user,
        )

    def test_create_transfer_request(self):
        transfer = self.create_pending_transfer()

        self.assertEqual(transfer.status, MemberTransferRequest.Status.PENDING)
        self.assertEqual(transfer.from_assembly, self.from_assembly)
        self.assertEqual(transfer.to_assembly, self.to_assembly)
        self.assertEqual(transfer.requested_by, self.source_user)
        self.assertEqual(self.member.assembly, self.from_assembly)
        self.assertTrue(
            AssemblyMembership.objects.filter(
                member=self.member,
                assembly=self.from_assembly,
                status=AssemblyMembershipStatus.ACTIVE,
                ended_on__isnull=True,
            ).exists()
        )

    def test_prevents_transfer_to_same_assembly(self):
        with self.assertRaises(ValidationError):
            create_transfer_request(
                member=self.member,
                to_assembly=self.from_assembly,
                effective_date=date(2026, 7, 10),
                requested_by=self.source_user,
            )

    def test_prevents_duplicate_pending_transfers(self):
        self.create_pending_transfer()

        with self.assertRaises(ValidationError):
            create_transfer_request(
                member=self.member,
                to_assembly=self.to_assembly,
                effective_date=date(2026, 7, 10),
                requested_by=self.source_user,
            )

    def test_accepting_transfer_moves_membership_only(self):
        transfer = self.create_pending_transfer()

        completed = accept_transfer_request(
            transfer=transfer,
            reviewed_by=self.receiving_user,
            notes="Accepted by receiving assembly",
        )

        self.member.refresh_from_db()
        self.assertEqual(completed.status, MemberTransferRequest.Status.COMPLETED)
        self.assertEqual(self.member.assembly, self.to_assembly)
        self.assertTrue(
            AssemblyMembership.objects.filter(
                member=self.member,
                assembly=self.from_assembly,
                status=AssemblyMembershipStatus.ENDED,
                end_reason=MembershipEndReason.TRANSFERRED,
                ended_on=date(2026, 7, 10),
                transfer=transfer,
            ).exists()
        )
        self.assertTrue(
            AssemblyMembership.objects.filter(
                member=self.member,
                assembly=self.to_assembly,
                status=AssemblyMembershipStatus.ACTIVE,
                joined_on=date(2026, 7, 10),
                ended_on__isnull=True,
            ).exists()
        )
        self.assertEqual(
            AssemblyMembership.objects.filter(
                member=self.member,
                status=AssemblyMembershipStatus.ACTIVE,
                ended_on__isnull=True,
            ).count(),
            1,
        )

    def test_rejecting_transfer_does_not_move_member(self):
        transfer = self.create_pending_transfer()

        rejected = reject_transfer_request(
            transfer=transfer,
            reviewed_by=self.receiving_user,
            rejection_reason="Member is not known to this assembly",
        )

        self.member.refresh_from_db()
        self.assertEqual(rejected.status, MemberTransferRequest.Status.REJECTED)
        self.assertEqual(self.member.assembly, self.from_assembly)
        self.assertFalse(
            AssemblyMembership.objects.filter(
                member=self.member,
                assembly=self.to_assembly,
            ).exists()
        )

    def test_cancelling_transfer_does_not_move_member(self):
        transfer = self.create_pending_transfer()

        cancelled = cancel_transfer_request(
            transfer=transfer,
            cancelled_by=self.source_user,
            notes="Created by mistake",
        )

        self.member.refresh_from_db()
        self.assertEqual(cancelled.status, MemberTransferRequest.Status.CANCELLED)
        self.assertEqual(self.member.assembly, self.from_assembly)
        self.assertFalse(
            AssemblyMembership.objects.filter(
                member=self.member,
                assembly=self.to_assembly,
            ).exists()
        )

    def test_old_tithe_and_report_records_are_not_migrated_on_accept(self):
        tithe = Tithe.objects.create(
            member=self.member,
            assembly=self.from_assembly,
            amount=Decimal("25.00"),
            payment_method="Cash",
            timestamp=date(2026, 1, 15),
        )
        old_report = tithe.report
        transfer = self.create_pending_transfer()

        accept_transfer_request(
            transfer=transfer,
            reviewed_by=self.receiving_user,
        )

        tithe.refresh_from_db()
        old_report.refresh_from_db()
        self.assertEqual(tithe.assembly, self.from_assembly)
        self.assertEqual(tithe.report, old_report)
        self.assertEqual(old_report.assembly, self.from_assembly)

    def test_source_assembly_cannot_accept_incoming_transfer(self):
        transfer = self.create_pending_transfer()
        client = APIClient()
        client.force_authenticate(user=self.source_user)

        response = client.post(
            f"/api/v1/people/member-transfers/{transfer.id}/accept/",
            {"notes": "Source should not accept"},
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        transfer.refresh_from_db()
        self.assertEqual(transfer.status, MemberTransferRequest.Status.PENDING)

    def test_transfer_visibility_uses_each_users_active_assembly(self):
        transfer = self.create_pending_transfer()

        # Membership in another selectable assembly must not expand the active
        # scope. Only user.church is authoritative for ordinary assembly users.
        self.other_user.assemblies.add(self.from_assembly)

        clients = {}
        for label, user in {
            "a": self.source_user,
            "b": self.receiving_user,
            "c": self.other_user,
        }.items():
            client = APIClient()
            client.force_authenticate(user=user)
            clients[label] = client

        def ids_for(client, view):
            response = client.get(
                f"/api/v1/people/member-transfers/{view}/",
                {
                    "page_size": 100,
                    # Deliberately hostile scope parameters must be ignored.
                    "from_assembly": self.from_assembly.id,
                    "to_assembly": self.to_assembly.id,
                },
            )
            self.assertEqual(response.status_code, 200)
            return {row["id"] for row in response.data["results"]}

        self.assertIn(transfer.id, ids_for(clients["a"], "outgoing"))
        self.assertNotIn(transfer.id, ids_for(clients["a"], "incoming"))
        self.assertIn(transfer.id, ids_for(clients["b"], "incoming"))
        self.assertNotIn(transfer.id, ids_for(clients["b"], "outgoing"))

        for view in ("incoming", "outgoing", "history"):
            self.assertNotIn(transfer.id, ids_for(clients["c"], view))

        rejected = clients["b"].post(
            f"/api/v1/people/member-transfers/{transfer.id}/reject/",
            {"rejection_reason": "Transfer declined"},
            format="json",
        )
        self.assertEqual(rejected.status_code, 200)

        self.assertIn(transfer.id, ids_for(clients["a"], "history"))
        self.assertIn(transfer.id, ids_for(clients["b"], "history"))
        self.assertNotIn(transfer.id, ids_for(clients["c"], "history"))

        forced_closed_in_open_view = clients["a"].get(
            "/api/v1/people/member-transfers/outgoing/",
            {"status": MemberTransferRequest.Status.REJECTED, "page_size": 100},
        )
        self.assertEqual(forced_closed_in_open_view.status_code, 200)
        self.assertNotIn(
            transfer.id,
            {row["id"] for row in forced_closed_in_open_view.data["results"]},
        )

        self.assertEqual(
            clients["c"].get(f"/api/v1/people/member-transfers/{transfer.id}/").status_code,
            404,
        )
        self.assertEqual(
            clients["c"].post(f"/api/v1/people/member-transfers/{transfer.id}/accept/").status_code,
            404,
        )
        self.assertEqual(
            clients["c"].post(
                f"/api/v1/people/member-transfers/{transfer.id}/reject/",
                {"rejection_reason": "Unauthorized"},
                format="json",
            ).status_code,
            404,
        )
        self.assertEqual(
            clients["c"].post(f"/api/v1/people/member-transfers/{transfer.id}/cancel/").status_code,
            404,
        )
        self.assertEqual(
            clients["c"].post(
                "/api/v1/people/member-transfers/",
                {
                    "member": self.member.id,
                    "to_assembly": self.external_assembly.id,
                    "effective_date": "2026-07-20",
                },
                format="json",
            ).status_code,
            403,
        )

    def test_transfer_lists_are_direction_scoped_for_assembly_users(self):
        outgoing_transfer = self.create_pending_transfer()
        incoming_member = self.create_member(self.other_assembly, "Incoming")
        incoming_transfer = create_transfer_request(
            member=incoming_member,
            to_assembly=self.from_assembly,
            effective_date=date(2026, 7, 11),
            reason="Moved back",
            requested_by=self.other_user,
        )
        exclusive_member = self.create_member(self.other_assembly, "Exclusive")
        exclusive_transfer = MemberTransferRequest.objects.create(
            member=exclusive_member,
            from_assembly=self.other_assembly,
            to_assembly=self.external_assembly,
            effective_date=date(2026, 7, 12),
            requested_by=self.other_user,
        )

        client = APIClient()
        client.force_authenticate(user=self.source_user)

        incoming = client.get(
            "/api/v1/people/member-transfers/incoming/",
            {"page_size": 100, "from_assembly": self.other_assembly.id},
        )
        outgoing = client.get(
            "/api/v1/people/member-transfers/outgoing/",
            {"page_size": 100, "to_assembly": self.external_assembly.id},
        )

        incoming_ids = {row["id"] for row in incoming.data["results"]}
        outgoing_ids = {row["id"] for row in outgoing.data["results"]}

        self.assertEqual(incoming.status_code, 200)
        self.assertEqual(outgoing.status_code, 200)
        self.assertIn(incoming_transfer.id, incoming_ids)
        self.assertNotIn(outgoing_transfer.id, incoming_ids)
        self.assertNotIn(exclusive_transfer.id, incoming_ids)
        self.assertIn(outgoing_transfer.id, outgoing_ids)
        self.assertNotIn(incoming_transfer.id, outgoing_ids)
        self.assertNotIn(exclusive_transfer.id, outgoing_ids)

    def test_transfer_history_is_closed_and_involves_user_assembly(self):
        history_member = self.create_member(self.from_assembly, "History")
        history_transfer = create_transfer_request(
            member=history_member,
            to_assembly=self.to_assembly,
            effective_date=date(2026, 7, 12),
            reason="Closed transfer",
            requested_by=self.source_user,
        )
        reject_transfer_request(
            transfer=history_transfer,
            reviewed_by=self.receiving_user,
            rejection_reason="Not moving",
        )
        open_transfer = self.create_pending_transfer()
        exclusive_member = self.create_member(self.other_assembly, "Other History")
        exclusive_transfer = MemberTransferRequest.objects.create(
            member=exclusive_member,
            from_assembly=self.other_assembly,
            to_assembly=self.external_assembly,
            status=MemberTransferRequest.Status.REJECTED,
            effective_date=date(2026, 7, 13),
            requested_by=self.other_user,
        )

        client = APIClient()
        client.force_authenticate(user=self.source_user)

        response = client.get(
            "/api/v1/people/member-transfers/history/",
            {"page_size": 100},
        )
        ids = {row["id"] for row in response.data["results"]}

        self.assertEqual(response.status_code, 200)
        self.assertIn(history_transfer.id, ids)
        self.assertNotIn(open_transfer.id, ids)
        self.assertNotIn(exclusive_transfer.id, ids)

    def test_direct_access_to_unrelated_transfer_is_denied(self):
        exclusive_member = self.create_member(self.other_assembly, "Unrelated")
        exclusive_transfer = MemberTransferRequest.objects.create(
            member=exclusive_member,
            from_assembly=self.other_assembly,
            to_assembly=self.external_assembly,
            effective_date=date(2026, 7, 14),
            requested_by=self.other_user,
        )

        client = APIClient()
        client.force_authenticate(user=self.source_user)

        detail = client.get(f"/api/v1/people/member-transfers/{exclusive_transfer.id}/")
        update = client.patch(
            f"/api/v1/people/member-transfers/{exclusive_transfer.id}/",
            {"notes": "Should not update"},
            format="json",
        )
        accept = client.post(f"/api/v1/people/member-transfers/{exclusive_transfer.id}/accept/")
        reject = client.post(
            f"/api/v1/people/member-transfers/{exclusive_transfer.id}/reject/",
            {"rejection_reason": "Should not reject"},
            format="json",
        )
        cancel = client.post(f"/api/v1/people/member-transfers/{exclusive_transfer.id}/cancel/")

        self.assertEqual(detail.status_code, 404)
        self.assertEqual(update.status_code, 404)
        self.assertEqual(accept.status_code, 404)
        self.assertEqual(reject.status_code, 404)
        self.assertEqual(cancel.status_code, 404)

    def test_regional_staff_retain_transfer_access_for_assigned_region(self):
        transfer = self.create_pending_transfer()
        exclusive_member = self.create_member(self.other_assembly, "Regional Hidden")
        exclusive_transfer = MemberTransferRequest.objects.create(
            member=exclusive_member,
            from_assembly=self.other_assembly,
            to_assembly=self.external_assembly,
            effective_date=date(2026, 7, 15),
            requested_by=self.other_user,
        )

        client = APIClient()
        client.force_authenticate(user=self.regional_user)

        detail = client.get(f"/api/v1/people/member-transfers/{transfer.id}/")
        outgoing = client.get(
            "/api/v1/people/member-transfers/outgoing/",
            {"page_size": 100},
        )
        exclusive_detail = client.get(f"/api/v1/people/member-transfers/{exclusive_transfer.id}/")

        outgoing_ids = {row["id"] for row in outgoing.data["results"]}

        self.assertEqual(detail.status_code, 200)
        self.assertEqual(outgoing.status_code, 200)
        self.assertIn(transfer.id, outgoing_ids)
        self.assertEqual(exclusive_detail.status_code, 404)

    def test_member_directory_exposes_pending_transfer_state(self):
        transfer = self.create_pending_transfer()
        client = APIClient()
        client.force_authenticate(user=self.source_user)

        response = client.get(
            "/api/v1/people/members/",
            {"page_size": 100},
        )

        self.assertEqual(response.status_code, 200)

        member_row = next(
            row
            for row in response.data["results"]
            if row["id"] == self.member.id
        )

        self.assertTrue(member_row["has_pending_transfer"])
        self.assertEqual(member_row["pending_transfer_id"], transfer.id)


class MembershipHouseholdWorkflowTests(TestCase):
    def setUp(self):
        region = Region.objects.create(name="Membership Region", code="MR")
        zone = Zone.objects.create(region=region, name="Membership Zone")
        self.assembly = Church.objects.create(name="Membership Assembly", zone=zone)
        self.other_assembly = Church.objects.create(name="Other Membership Assembly", zone=zone)
        self.user = User.objects.create_user(
            first_name="Assembly", last_name="User", email="assembly-user@example.com",
            username="assembly-user", password="password", church=self.assembly,
        )
        self.other_user = User.objects.create_user(
            first_name="Other", last_name="User", email="other-user-2@example.com",
            username="other-user-2", password="password", church=self.other_assembly,
        )
        self.member = self.make_member(self.assembly, "History")

    def make_member(self, assembly, first_name):
        return Member.objects.create(
            assembly=assembly, first_name=first_name, last_name="Person",
            date_of_birth=date(1992, 2, 2), gender="Male", country="Botswana",
            membersince=date(2020, 1, 1),
        )

    def make_membership(self, member=None, **overrides):
        values = {
            "member": member or self.member,
            "assembly": (member or self.member).assembly,
            "joined_on": date(2020, 1, 1),
            "status": AssemblyMembershipStatus.ACTIVE,
        }
        values.update(overrides)
        return AssemblyMembership.objects.create(**values)

    def test_only_one_assembly_membership_model_is_registered(self):
        registered = [
            model for model in django_apps.get_app_config("people").get_models()
            if model._meta.model_name == "assemblymembership"
        ]
        self.assertEqual(registered, [AssemblyMembership])

    def test_current_and_historical_membership_invariants(self):
        current = self.make_membership()
        with self.assertRaises(DjangoValidationError):
            self.make_membership(
                assembly=self.other_assembly,
                status=AssemblyMembershipStatus.INACTIVE,
            )

        current.status = AssemblyMembershipStatus.ENDED
        current.ended_on = date(2022, 1, 1)
        current.end_reason = MembershipEndReason.RESIGNED
        current.save()
        second_history = self.make_membership(
            assembly=self.other_assembly,
            joined_on=date(2023, 1, 1),
            ended_on=date(2024, 1, 1),
            end_reason=MembershipEndReason.RELOCATED,
            status=AssemblyMembershipStatus.ENDED,
        )
        self.assertEqual(AssemblyMembership.objects.former().filter(member=self.member).count(), 2)
        self.assertTrue(second_history.is_former)

    def test_membership_date_reason_and_transfer_validation(self):
        invalid = AssemblyMembership(
            member=self.member, assembly=self.assembly,
            joined_on=date(2024, 1, 2), status=AssemblyMembershipStatus.ENDED,
        )
        with self.assertRaises(DjangoValidationError):
            invalid.full_clean()
        invalid.ended_on = date(2024, 1, 1)
        invalid.end_reason = MembershipEndReason.RESIGNED
        with self.assertRaises(DjangoValidationError):
            invalid.full_clean()
        invalid.ended_on = date(2024, 1, 3)
        invalid.end_reason = MembershipEndReason.TRANSFERRED
        with self.assertRaises(DjangoValidationError):
            invalid.full_clean()

    def test_household_history_and_primary_contact_rules(self):
        membership = self.make_membership()
        household = Household.objects.create(assembly=self.assembly, name="Person Household")
        household_member = HouseholdMember.objects.create(
            household=household, member=self.member, role=HouseholdRole.HEAD,
            is_primary_contact=True, joined_on=date(2020, 1, 1),
        )
        with self.assertRaises(DjangoValidationError):
            HouseholdMember.objects.create(
                household=Household.objects.create(assembly=self.assembly, name="Second Household"),
                member=self.member, joined_on=date(2021, 1, 1),
            )
        household_member.leave(left_on=date(2022, 1, 1), updated_by=self.user)
        household_member.refresh_from_db()
        self.assertEqual(household_member.left_on, date(2022, 1, 1))
        self.assertFalse(household_member.is_primary_contact)
        self.assertTrue(AssemblyMembership.objects.filter(pk=membership.pk).exists())

    def test_closed_or_cross_assembly_household_rejects_member(self):
        self.make_membership()
        closed = Household.objects.create(assembly=self.assembly, name="Closed Household")
        closed.close(updated_by=self.user)
        with self.assertRaises(DjangoValidationError):
            HouseholdMember.objects.create(household=closed, member=self.member)
        other_household = Household.objects.create(assembly=self.other_assembly, name="Remote Household")
        with self.assertRaises(DjangoValidationError):
            HouseholdMember.objects.create(household=other_household, member=self.member)

    def test_transfer_completion_is_idempotent_and_ends_household_membership(self):
        source = self.make_membership()
        household = Household.objects.create(assembly=self.assembly, name="Transfer Household")
        household_member = HouseholdMember.objects.create(
            household=household, member=self.member, joined_on=date(2020, 1, 1)
        )
        transfer = create_transfer_request(
            member=self.member, to_assembly=self.other_assembly,
            effective_date=date(2026, 1, 1), requested_by=self.user,
        )
        completed, destination = complete_transfer_request(
            transfer=transfer, completed_by=self.other_user,
        )
        completed_again, destination_again = complete_transfer_request(
            transfer=completed, completed_by=self.other_user,
        )
        source.refresh_from_db()
        household_member.refresh_from_db()
        self.member.refresh_from_db()
        self.assertEqual(source.end_reason, MembershipEndReason.TRANSFERRED)
        self.assertEqual(source.transfer, transfer)
        self.assertEqual(household_member.left_on, date(2026, 1, 1))
        self.assertEqual(destination.pk, destination_again.pk)
        self.assertEqual(self.member.assembly, self.other_assembly)
        self.assertEqual(AssemblyMembership.objects.current().filter(member=self.member).count(), 1)

    def test_former_member_api_is_assembly_scoped_and_readmission_preserves_history(self):
        historical = self.make_membership(
            status=AssemblyMembershipStatus.ENDED,
            ended_on=date(2024, 1, 1), end_reason=MembershipEndReason.RESIGNED,
        )
        self.member.assembly = None
        self.member.save(update_fields=["assembly", "updated_at"])
        client = APIClient()
        client.force_authenticate(user=self.user)
        response = client.get("/api/v1/people/former-members/", {"reason": "resigned"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row["id"] for row in response.data["results"]], [historical.pk])
        readmit = client.post(
            f"/api/v1/people/former-members/{historical.pk}/readmit/",
            {"joined_on": "2026-01-02"}, format="json",
        )
        self.assertEqual(readmit.status_code, 201)
        historical.refresh_from_db()
        self.assertEqual(historical.status, AssemblyMembershipStatus.ENDED)
        self.assertEqual(AssemblyMembership.objects.current().filter(member=self.member).count(), 1)

    def test_household_api_uses_persisted_scoped_data(self):
        self.make_membership()
        client = APIClient()
        client.force_authenticate(user=self.user)
        created = client.post(
            "/api/v1/people/households/", {"name": "API Household"}, format="json"
        )
        self.assertEqual(created.status_code, 201)
        added = client.post(
            f"/api/v1/people/households/{created.data['id']}/add-member/",
            {"member": self.member.pk, "joined_on": "2020-01-01", "role": "head"},
            format="json",
        )
        self.assertEqual(added.status_code, 201)
        self.assertTrue(HouseholdMember.objects.filter(pk=added.data["id"]).exists())
        other = Household.objects.create(assembly=self.other_assembly, name="Hidden Household")
        listing = client.get("/api/v1/people/households/", {"page_size": 100})
        ids = {row["id"] for row in listing.data["results"]}
        self.assertIn(created.data["id"], ids)
        self.assertNotIn(other.pk, ids)
